#!/usr/bin/env python3
"""Exporta TODOS os dados brutos do levantamento em dois formatos.

  planilha : output/epidemio/exports/HICD_cardiopatias_dados_brutos.xlsx
             uma aba por entidade + a aba `evolucoes` com as 130.636 linhas.
  docx     : output/epidemio/exports/HICD_evolucoes_volNN_*.docx
             o texto integral das evoluções, um capítulo por paciente.

Por que o docx sai em volumes: são ~347 MB de texto. Nenhum processador de
texto abre isso num arquivo só — nem o .docx é o gargalo (é zip), é o Word
montando o layout de milhões de parágrafos. Os volumes são cortados por
paciente, nunca no meio de um prontuário.

O .docx é escrito direto como zip, em streaming. python-docx montaria a árvore
XML inteira em memória (dezenas de GB) e não terminaria.

Uso:
  python3 scripts/exportar_brutos.py [--planilha] [--docx] [--mb-por-volume 25]
"""
import argparse, csv, glob, json, os, re, zipfile
from datetime import datetime

csv.field_size_limit(10 ** 9)
BASE = 'output/epidemio'
OUT = os.path.join(BASE, 'exports')
LIMITE_CELULA = 32767          # limite duro do formato xlsx, não escolha nossa
INVALIDO_XML = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def limpar(t):
    """Texto de prontuário para XML: <br> vira quebra, tags somem, controles somem."""
    t = re.sub(r'<br\s*/?>', '\n', str(t or ''), flags=re.I)
    t = re.sub(r'<[^>]+>', '', t)
    t = (t.replace('&amp;', '&').replace('&lt;', '<').replace('&gt;', '>')
          .replace('&nbsp;', ' ').replace('&quot;', '"').replace('&#39;', "'"))
    return INVALIDO_XML.sub('', t).replace('\xa0', ' ')


def prontuarios():
    """Todos os raw_*.json, em ordem numérica de prontuário."""
    fs = glob.glob(os.path.join(BASE, 'raw_*.json'))
    return sorted(fs, key=lambda f: int(os.path.basename(f)[4:-5]))


def ler_csv(caminho):
    if not os.path.exists(caminho):
        return [], []
    with open(caminho, encoding='utf-8') as fh:
        r = list(csv.DictReader(fh))
    return (list(r[0].keys()) if r else []), r


# --------------------------------------------------------------------------- #
# Planilha
# --------------------------------------------------------------------------- #
def planilha(destino):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)   # streaming: não guarda as linhas em memória

    def aba(nome, colunas, linhas, larguras=None):
        ws = wb.create_sheet(nome)
        if larguras:
            ws.column_dimensions  # força criação
            for i, w in enumerate(larguras, start=1):
                ws.column_dimensions[get_column_letter(i)].width = w
        ws.append(colunas)
        n = 0
        for l in linhas:
            ws.append([l.get(c) if isinstance(l, dict) else l[i] for i, c in enumerate(colunas)])
            n += 1
        print(f'  aba {nome}: {n} linhas', flush=True)
        return n

    # 1. dicionário -------------------------------------------------------- #
    dic = [
        ('dicionario', '—', 'esta aba: o que cada aba e cada coluna significam'),
        ('pacientes', 'output/epidemio/dataset_pacientes.csv',
         'um paciente por linha. Demografia do cadastro e do SAME, internação-índice, '
         'classificação cardiológica e desfecho agregado. Duplicatas de prontuário já fundidas.'),
        ('internacoes', 'output/epidemio/dataset_internacoes.csv',
         'uma internação por linha. Entrada, saída, setor de alta, setor de entrada derivado '
         'das evoluções, trajetória, boletim de emergência associado e desfecho resolvido.'),
        ('classificacao', 'docs/classificacao_cardiologica.csv',
         'a classificação cardiológica de cada prontuário, com a justificativa escrita e a '
         'procedência (lida à mão × triagem automática).'),
        ('evolucoes', 'raw_*.json → campo evolucoes',
         'TODAS as evoluções coletadas, texto integral. É a fonte primária de tudo: '
         'óbito, desfecho, trajetória por setor e diagnóstico cardiológico saem daqui.'),
        ('boletins_emergencia', 'be_*.json',
         'Boletim de Emergência: a porta de entrada. Motivo da entrada, CID e, quando '
         'preenchida, a classificação de risco de Manchester.'),
        ('cadastro_same', 'raw_*.json → campo cadastroSame',
         'cadastro do SAME: raça/cor, deficiência, CNS e código IBGE do município. '
         'ATENÇÃO: cobertura de raça/cor é baixa — meça antes de analisar.'),
        ('marco_amostral', 'output/epidemio/marco_amostral.json',
         'quem entrou na amostra e por quê: parecer de cardiologista ou varredura por termo.'),
        ('avisos', '—', 'limitações que precisam acompanhar qualquer número tirado daqui'),
    ]
    aba('dicionario', ['aba', 'origem', 'o que é'], dic, [22, 46, 110])

    avisos = [
        ('desfecho', '22,9% das internações ficam "Sem informação". O HICD não obriga resumo de alta.'),
        ('óbito', 'Não existe campo de óbito. Ele é detectado no TEXTO das evoluções — o módulo '
                  'RALTA fica vazio justamente nos óbitos, porque não se escreve resumo de alta para quem morreu.'),
        ('setor', 'A coluna setor_alta vem do módulo Inter e é o setor de ALTA, não o de entrada. '
                  'O percurso real está em trajetoria_na_internacao, derivado das evoluções.'),
        ('raça/cor', 'Cobertura baixa (~38%). Declarar como limitação, não analisar.'),
        ('classificação de risco', 'Preenchida em ~0% das congênitas. Inutilizável.'),
        ('módulo Inter', 'Devolve HTTP 500 em prontuários muito longos. 6 prontuários ficaram sem '
                         'dados de internação: 20969, 33549, 34387, 40600, 41522, 5641.'),
        ('permanência', 'Há registros corrompidos na origem (saída anterior à entrada, internação de '
                        '553 dias). A coluna permanencia_dias vem vazia nesses casos e o motivo fica em permanencia_invalida.'),
        ('FOP isolado', 'Fica no numerador por decisão do orientando, mas 19,5% das congênitas são FOP '
                        'isolado e NÃO há corte de idade. Toda tabela precisa da coluna "sem FOP isolado".'),
        ('classificação automática', '297 prontuários foram classificados só pela triagem automática. '
                                     'Erro residual medido: 1,3%–1,7%. "A triagem não achou" nunca significa "não existe".'),
        ('texto truncado', f'O formato xlsx limita cada célula a {LIMITE_CELULA:,} caracteres. '
                           '91 evoluções (0,07%) foram cortadas — a coluna texto_truncado marca quais. '
                           'O texto íntegro dessas está nos arquivos .docx.'),
        ('dados identificáveis', 'Esta planilha contém nome, data de nascimento e prontuário de '
                                 'pacientes pediátricos. Tratar conforme o protocolo do CEP.'),
    ]
    aba('avisos', ['tema', 'aviso'], avisos, [26, 130])

    # 2. abas derivadas dos CSVs já prontos -------------------------------- #
    for nome, caminho in [('pacientes', f'{BASE}/dataset_pacientes.csv'),
                          ('internacoes', f'{BASE}/dataset_internacoes.csv'),
                          ('classificacao', 'docs/classificacao_cardiologica.csv')]:
        cols, linhas = ler_csv(caminho)
        if cols:
            aba(nome, cols, linhas)

    # 3. cadastro do SAME + marco ------------------------------------------ #
    same_cols = ['prontuario', 'nome', 'sexo', 'dataNascimento', 'racaCor', 'racaCorCodigo',
                 'deficiencia', 'nacionalidade', 'municipio', 'municipioIbge', 'uf', 'bairro',
                 'cns', 'cpf', 'nomeMae', 'total_evolucoes', 'total_internacoes', 'falha_inter']
    same_linhas = []
    for f in prontuarios():
        d = json.load(open(f, encoding='utf-8'))
        b = (d.get('cadastro') or {}).get('dadosBasicos') or {}
        s = d.get('cadastroSame') or {}
        ev, it = d.get('evolucoes'), d.get('internacoes')
        same_linhas.append({
            'prontuario': d.get('prontuario'), 'nome': b.get('nome') or s.get('nome'),
            'sexo': b.get('sexo') or s.get('sexo'),
            'dataNascimento': str(b.get('dataNascimento') or s.get('dataNascimento') or '')[:10],
            'racaCor': s.get('racaCor'), 'racaCorCodigo': s.get('racaCorCodigo'),
            'deficiencia': s.get('deficiencia'), 'nacionalidade': s.get('nacionalidade'),
            'municipio': ((d.get('cadastro') or {}).get('endereco') or {}).get('municipio'),
            'municipioIbge': s.get('municipioIbge'), 'uf': s.get('uf'), 'bairro': s.get('bairro'),
            'cns': s.get('cns'), 'cpf': s.get('cpf'), 'nomeMae': s.get('nomeMae') or b.get('nomeMae'),
            'total_evolucoes': len(ev) if isinstance(ev, list) else None,
            'total_internacoes': len(it) if isinstance(it, list) else None,
            'falha_inter': 'Sim' if not isinstance(it, list) else 'Não',
        })
    aba('cadastro_same', same_cols, same_linhas)

    marco_f = f'{BASE}/marco_amostral.json'
    if os.path.exists(marco_f):
        m = json.load(open(marco_f, encoding='utf-8'))
        aba('marco_amostral', ['prontuario', 'fonte', 'n_avaliacoes_cardio', 'primeira_data',
                               'ultima_data', 'profissionais'],
            [{'prontuario': k, 'fonte': v.get('fonte'), 'n_avaliacoes_cardio': len(v.get('datas') or []),
              'primeira_data': (v.get('datas') or [None])[0],
              'ultima_data': (v.get('datas') or [None])[-1],
              'profissionais': '; '.join(v.get('profissionais') or [])}
             for k, v in sorted(m.items(), key=lambda x: int(x[0]))])

    # 4. boletins de emergência -------------------------------------------- #
    be_linhas = []
    for f in sorted(glob.glob(os.path.join(BASE, 'be_*.json'))):
        d = json.load(open(f, encoding='utf-8'))
        for b in d.get('boletins') or []:
            t = b.get('triagem') or {}
            be_linhas.append({
                'prontuario': d.get('prontuario'), 'be': b.get('be'), 'entrada': b.get('entrada'),
                'saida': b.get('saida'), 'motivo': b.get('motivo'), 'cid': b.get('cid'),
                'classificacao_risco': t.get('classificacaoRisco'), 'queixa': t.get('queixa'),
                'pa': t.get('pa'), 'fc': t.get('fc'), 'fr': t.get('fr'),
                'temperatura': t.get('temperatura'), 'saturacao': t.get('saturacao'),
                'peso': t.get('peso'), 'glicemia': t.get('glicemia'),
            })
    aba('boletins_emergencia',
        ['prontuario', 'be', 'entrada', 'saida', 'motivo', 'cid', 'classificacao_risco',
         'queixa', 'pa', 'fc', 'fr', 'temperatura', 'saturacao', 'peso', 'glicemia'], be_linhas)

    # 5. evoluções — a aba grande ------------------------------------------ #
    cols = ['prontuario', 'nome', 'data', 'profissional', 'atividade', 'clinica_leito',
            'hipoteses_diagnosticas', 'diagnosticos_anteriores', 'n_caracteres',
            'texto_truncado', 'descricao']
    ws = wb.create_sheet('evolucoes')
    for i, w in enumerate([12, 34, 18, 30, 22, 26, 46, 46, 13, 14, 120], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.append(cols)
    n = cortadas = 0
    for f in prontuarios():
        d = json.load(open(f, encoding='utf-8'))
        pront = d.get('prontuario')
        nome = ((d.get('cadastro') or {}).get('dadosBasicos') or {}).get('nome')
        ev = d.get('evolucoes')
        if not isinstance(ev, list):
            continue
        for e in ev:
            t = limpar(e.get('descricao'))
            corta = len(t) > LIMITE_CELULA
            if corta:
                cortadas += 1
            ws.append([pront, nome, e.get('data'), e.get('profissional'), e.get('atividade'),
                       e.get('clinicaLeito'),
                       ' | '.join(e.get('hipoteses') or [])[:LIMITE_CELULA],
                       ' | '.join(e.get('anteriores') or [])[:LIMITE_CELULA],
                       len(t), 'Sim' if corta else 'Não',
                       t[:LIMITE_CELULA]])
            n += 1
            if n % 20000 == 0:
                print(f'    evolucoes: {n:,}...', flush=True)
    print(f'  aba evolucoes: {n:,} linhas ({cortadas} truncadas)', flush=True)

    os.makedirs(os.path.dirname(destino), exist_ok=True)
    wb.save(destino)
    mb = os.path.getsize(destino) / 1e6
    print(f'\nPLANILHA: {destino}  ({mb:.1f} MB)')


# --------------------------------------------------------------------------- #
# DOCX em streaming
# --------------------------------------------------------------------------- #
CONTENT_TYPES = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>
</Types>'''

RELS = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>'''

DOC_RELS = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>
</Relationships>'''


def _estilo(sid, nome, tam, cor, bold, outline=None, espaco_antes=240):
    ol = f'<w:outlineLvl w:val="{outline}"/>' if outline is not None else ''
    return (f'<w:style w:type="paragraph" w:styleId="{sid}"><w:name w:val="{nome}"/>'
            f'<w:basedOn w:val="Normal"/><w:pPr><w:keepNext/>'
            f'<w:spacing w:before="{espaco_antes}" w:after="120"/>{ol}</w:pPr>'
            f'<w:rPr><w:b w:val="{"1" if bold else "0"}"/><w:color w:val="{cor}"/>'
            f'<w:sz w:val="{tam}"/></w:rPr></w:style>')


STYLES = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
          '<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
          '<w:docDefaults><w:rPrDefault><w:rPr>'
          '<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/><w:sz w:val="20"/>'
          '</w:rPr></w:rPrDefault></w:docDefaults>'
          '<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/>'
          '<w:pPr><w:spacing w:after="60"/></w:pPr></w:style>'
          + _estilo('Titulo', 'Title', 44, '1F3864', True, espaco_antes=0)
          + _estilo('Heading1', 'heading 1', 32, '1F3864', True, 0, 480)
          + _estilo('Heading2', 'heading 2', 24, '2E5A88', True, 1, 320)
          + _estilo('Heading3', 'heading 3', 21, '555555', True, 2, 200)
          + '<w:style w:type="paragraph" w:styleId="Meta"><w:name w:val="Meta"/>'
            '<w:basedOn w:val="Normal"/><w:rPr><w:i/><w:color w:val="666666"/>'
            '<w:sz w:val="17"/></w:rPr></w:style>'
          + '<w:style w:type="paragraph" w:styleId="Corpo"><w:name w:val="Corpo"/>'
            '<w:basedOn w:val="Normal"/><w:pPr><w:spacing w:after="0"/></w:pPr></w:style>'
          + '</w:styles>')

DOC_ABRE = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:body>')
DOC_FECHA = ('<w:sectPr><w:pgSz w:w="11906" w:h="16838"/>'
             '<w:pgMar w:top="1134" w:right="1134" w:bottom="1134" w:left="1134"/>'
             '</w:sectPr></w:body></w:document>')


def esc(t):
    return (t.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;'))


def par(texto, estilo='Corpo'):
    """Um parágrafo. Quebras internas viram <w:br/> em vez de parágrafos novos:
    130 mil evoluções × 40 linhas cada estouraria o limite prático do Word."""
    partes = esc(texto).split('\n')
    runs = '<w:br/>'.join(f'<w:t xml:space="preserve">{p}</w:t>' for p in partes)
    return (f'<w:p><w:pPr><w:pStyle w:val="{estilo}"/></w:pPr>'
            f'<w:r>{runs}</w:r></w:p>')


def quebra_pagina():
    return '<w:p><w:r><w:br w:type="page"/></w:r></w:p>'


def escrever_docx(caminho, blocos):
    """blocos: iterável de strings de XML de parágrafo já prontas."""
    with zipfile.ZipFile(caminho, 'w', zipfile.ZIP_DEFLATED, compresslevel=6) as z:
        z.writestr('[Content_Types].xml', CONTENT_TYPES)
        z.writestr('_rels/.rels', RELS)
        z.writestr('word/_rels/document.xml.rels', DOC_RELS)
        z.writestr('word/styles.xml', STYLES)
        with z.open('word/document.xml', 'w') as fh:
            fh.write(DOC_ABRE.encode('utf-8'))
            for b in blocos:
                fh.write(b.encode('utf-8'))
            fh.write(DOC_FECHA.encode('utf-8'))


def docx(mb_por_volume):
    os.makedirs(OUT, exist_ok=True)
    limite = int(mb_por_volume * 1_000_000)

    # classificação e nomes, para o cabeçalho de cada capítulo
    _, cls = ler_csv('docs/classificacao_cardiologica.csv')
    CLS = {r['prontuario']: r for r in cls}

    arquivos = prontuarios()
    # 1ª passada: medir cada paciente e fechar os volumes sem partir prontuário
    tam = {}
    for f in arquivos:
        d = json.load(open(f, encoding='utf-8'))
        ev = d.get('evolucoes')
        tam[f] = sum(len(e.get('descricao') or '') for e in ev) if isinstance(ev, list) else 0

    volumes, atual, acc = [], [], 0
    for f in arquivos:
        if atual and acc + tam[f] > limite:
            volumes.append(atual); atual, acc = [], 0
        atual.append(f); acc += tam[f]
    if atual:
        volumes.append(atual)

    print(f'{len(arquivos)} prontuários · {sum(tam.values())/1e6:.0f} MB de texto '
          f'· {len(volumes)} volumes de até {mb_por_volume} MB\n', flush=True)

    indice = []
    for vi, vol in enumerate(volumes, start=1):
        p0 = os.path.basename(vol[0])[4:-5]
        p1 = os.path.basename(vol[-1])[4:-5]
        nome = f'HICD_evolucoes_vol{vi:02d}_prontuarios_{p0}_a_{p1}.docx'
        destino = os.path.join(OUT, nome)

        def blocos(vol=vol, vi=vi, p0=p0, p1=p1):
            yield par(f'Evoluções do prontuário eletrônico — volume {vi} de {len(volumes)}', 'Titulo')
            yield par(f'Prontuários {p0} a {p1} · {len(vol)} pacientes', 'Meta')
            yield par('Estudo epidemiológico de cardiopatias congênitas — HICD, 2024–2025. '
                      'Texto integral das evoluções, transcrito do prontuário sem edição. '
                      f'Gerado em {datetime.now().strftime("%d/%m/%Y %H:%M")}.', 'Meta')
            yield par('CONTÉM DADOS IDENTIFICÁVEIS DE PACIENTES PEDIÁTRICOS. '
                      'Tratar conforme o protocolo aprovado pelo CEP.', 'Meta')
            for f in vol:
                d = json.load(open(f, encoding='utf-8'))
                pront = str(d.get('prontuario'))
                b = (d.get('cadastro') or {}).get('dadosBasicos') or {}
                c = CLS.get(pront, {})
                ev = d.get('evolucoes')
                ev = ev if isinstance(ev, list) else []
                yield quebra_pagina()
                yield par(f'Prontuário {pront} — {b.get("nome") or "(sem nome no cadastro)"}', 'Heading1')
                meta = [f'Nascimento: {str(b.get("dataNascimento") or "")[:10] or "—"}',
                        f'Sexo: {b.get("sexo") or "—"}',
                        f'Evoluções: {len(ev)}']
                if c:
                    meta.append(f'Cardiopatia congênita: {c.get("cardiopatia_congenita")}')
                    if c.get('classe'):
                        meta.append(f'Classe: {c["classe"]}')
                    meta.append(f'Procedência da classificação: {c.get("origem_classificacao")}')
                yield par(' · '.join(meta), 'Meta')
                if c.get('dx_cardio'):
                    yield par(f'Diagnóstico cardiológico: {limpar(c["dx_cardio"])}', 'Meta')
                if not ev:
                    yield par('(nenhuma evolução coletada para este prontuário)', 'Meta')
                for e in sorted(ev, key=lambda x: str(x.get('data') or '')):
                    cab = f'{e.get("data") or "sem data"} — {e.get("profissional") or "sem profissional"}'
                    yield par(cab, 'Heading2')
                    sub = [x for x in [e.get('atividade'), e.get('clinicaLeito')] if x]
                    if sub:
                        yield par(' · '.join(sub), 'Meta')
                    t = limpar(e.get('descricao')).strip()
                    yield par(t if t else '(evolução sem texto)')

        escrever_docx(destino, blocos())
        mb = os.path.getsize(destino) / 1e6
        pacs = len(vol)
        evs = sum(len(json.load(open(f, encoding='utf-8')).get('evolucoes') or [])
                  for f in vol)
        indice.append((nome, p0, p1, pacs, evs, round(mb, 1)))
        print(f'  vol {vi:02d}: {nome}  ({pacs} pacientes, {evs:,} evoluções, {mb:.1f} MB)', flush=True)

    # índice dos volumes, em CSV (entra na planilha numa próxima rodada se quiser)
    with open(os.path.join(OUT, 'indice_volumes.csv'), 'w', encoding='utf-8', newline='') as fh:
        w = csv.writer(fh)
        w.writerow(['arquivo', 'prontuario_inicial', 'prontuario_final', 'pacientes', 'evolucoes', 'mb'])
        w.writerows(indice)
    print(f'\nÍNDICE: {os.path.join(OUT, "indice_volumes.csv")}')


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--planilha', action='store_true')
    ap.add_argument('--docx', action='store_true')
    ap.add_argument('--mb-por-volume', type=float, default=25)
    a = ap.parse_args()
    if not (a.planilha or a.docx):
        a.planilha = a.docx = True
    os.makedirs(OUT, exist_ok=True)
    if a.planilha:
        print('=== PLANILHA ===', flush=True)
        planilha(os.path.join(OUT, 'HICD_cardiopatias_dados_brutos.xlsx'))
    if a.docx:
        print('\n=== DOCX ===', flush=True)
        docx(a.mb_por_volume)
